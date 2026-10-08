"""Local source ingestion. Documents and public links are treated as data."""
from __future__ import annotations

import asyncio
import io
import ipaddress
import json
import math
import re
import shutil
import socket
import sys
import tempfile
import time
import uuid
import zipfile
from html import unescape
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse

import httpx
from fastapi import APIRouter, File, HTTPException, UploadFile
from pydantic import BaseModel

SOURCES_DIR = Path(__file__).resolve().parents[1] / "sources"
SOURCES_DIR.mkdir(exist_ok=True)
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_SOURCE_CHARS = 180_000
MAX_VIDEO_SECONDS = 7200
router = APIRouter()
_YOUTUBE_LOCK = asyncio.Semaphore(2)


class SourceError(ValueError):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


class YouTubeSource(BaseModel):
    url: str
    elevenlabs_key: str | None = None


class ArticleURLSource(BaseModel):
    url: str


def clean_text(text: str) -> str:
    text = unescape(text).replace("\x00", "").replace("\r\n", "\n")
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]*\n(?:[ \t]*\n)+", "\n\n", text)
    return text.strip()


class _ArticleHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self.title: list[str] = []
        self.hidden = 0
        self.in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self.hidden += 1
        if tag == "title":
            self.in_title = True
        if tag in ("p", "div", "br", "li", "h1", "h2", "h3"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript"):
            self.hidden = max(0, self.hidden - 1)
        if tag == "title":
            self.in_title = False
        if tag in ("p", "div", "li", "h1", "h2", "h3"):
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)
            if self.in_title:
                self.title.append(data)


def extract_article(data: bytes, filename: str) -> dict:
    """Extract supported formats in memory; never execute embedded content."""
    name = Path(filename.replace("\\", "/")).name or "Article"
    suffix = Path(name).suffix.lower()
    title = Path(name).stem.replace("_", " ").strip() or "Article"
    warnings: list[str] = []
    if len(data) > MAX_UPLOAD_BYTES:
        raise SourceError("Upload a document smaller than 20 MB.", 413)
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader
            if not data.lstrip().startswith(b"%PDF-"):
                raise SourceError("This file is not a valid PDF.")
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise SourceError("This PDF is password protected. Upload an unlocked copy.")
            if len(reader.pages) > 300:
                raise SourceError("Upload a paper or article of at most 300 pages, or a selected section.")
            parts = []
            total = 0
            for page in reader.pages:
                part = page.extract_text() or ""
                total += len(part)
                if total > MAX_SOURCE_CHARS:
                    raise SourceError("This paper is too long. Upload a selected section (up to 180,000 text characters).")
                parts.append(part)
            text = "\n\n".join(parts)
            metadata_title = getattr(reader.metadata, "title", None)
            if metadata_title and str(metadata_title).strip():
                title = str(metadata_title).strip()[:200]
            warnings.append("PDF diagrams and images are not read. Check equations and extracted text before generating.")
        elif suffix == ".docx":
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 80 * 1024 * 1024:
                    raise SourceError("This Word document expands to too much data. Upload a smaller section.")
                if "word/document.xml" not in archive.namelist():
                    raise SourceError("This file is not a valid Word document.")
            document = Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows)
            text = "\n".join(parts)
            warnings.append("Embedded images and diagrams are not read.")
        elif suffix in (".txt", ".md", ".html", ".htm"):
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise SourceError("Save this text file as UTF-8, then upload it again.")
            if suffix in (".html", ".htm"):
                from lxml import html
                document = html.fromstring(text)
                title = " ".join(document.xpath("//title/text()")).strip()[:200] or title
                for node in document.xpath("//script|//style|//noscript|//nav|//header|//footer|//form"):
                    node.drop_tree()
                candidates = document.xpath("//article|//main")
                content = max(candidates, key=lambda node: len(node.text_content())) if candidates else document
                parser = _ArticleHTML()
                parser.feed(html.tostring(content, encoding="unicode"))
                text = "".join(parser.parts)
        else:
            raise SourceError("Upload a PDF, Word (.docx), text (.txt), Markdown (.md), or HTML article.")
    except SourceError:
        raise
    except Exception as e:
        raise SourceError("The document could not be read. Check that it opens correctly, or upload a text version.") from e
    text = clean_text(text)
    if len(text) < 80:
        raise SourceError("The document has too little readable text. Scanned PDFs need OCR; upload a selectable-text PDF or text version.")
    if len(text) > MAX_SOURCE_CHARS:
        raise SourceError("This source is too long. Upload a selected section (up to 180,000 text characters).")
    return {"type": "article", "title": title[:200], "text": text, "filename": name, "url": None, "warnings": warnings}


def youtube_video_id(url: str) -> str:
    url = url.strip()
    if not url:
        raise SourceError("Paste a public YouTube video link.")
    if not re.match(r"^https?://", url, re.I):
        url = "https://" + url
    try:
        parsed = urlparse(url)
        host = (parsed.hostname or "").lower()
        port = parsed.port
    except ValueError as e:
        raise SourceError("Use a normal public YouTube video link.") from e
    if parsed.username or parsed.password or port not in (None, 80, 443):
        raise SourceError("Use a normal public YouTube video link.")
    if host == "youtu.be":
        candidate = parsed.path.strip("/")
    elif host in ("youtube.com", "www.youtube.com", "m.youtube.com", "music.youtube.com"):
        if parsed.path == "/watch":
            candidate = parse_qs(parsed.query).get("v", [""])[0]
        else:
            parts = parsed.path.strip("/").split("/")
            candidate = parts[1] if len(parts) == 2 and parts[0] in ("shorts", "embed", "live") else ""
    else:
        raise SourceError("Only YouTube video links are supported here.")
    if not re.fullmatch(r"[A-Za-z0-9_-]{11}", candidate):
        raise SourceError("This link does not identify a YouTube video. Use a watch, Shorts, or youtu.be link.")
    return candidate


def _fetch_captions(video_id: str) -> tuple[str, float, bool, str]:
    from requests import Session
    from youtube_transcript_api import YouTubeTranscriptApi
    # Bound individual network calls even when the upstream library retries.
    class TimeoutSession(Session):
        def request(self, *args, **kwargs):
            kwargs.setdefault("timeout", 20)
            return super().request(*args, **kwargs)
    transcripts = list(YouTubeTranscriptApi(http_client=TimeoutSession()).list(video_id))
    if not transcripts:
        raise SourceError("No captions are available.")
    transcripts.sort(key=lambda t: (not str(t.language_code).startswith("en"), t.is_generated))
    fetched = transcripts[0].fetch()
    text = clean_text(" ".join(snippet.text for snippet in fetched))
    duration = max((snippet.start + snippet.duration for snippet in fetched), default=0)
    return text, duration, fetched.is_generated, fetched.language


async def _youtube_title(canonical_url: str, video_id: str) -> str:
    try:
        async with httpx.AsyncClient(timeout=15) as client:
            response = await client.get("https://www.youtube.com/oembed", params={"url": canonical_url, "format": "json"})
            if response.is_success:
                metadata = response.json()
                if isinstance(metadata, dict):
                    return str(metadata.get("title") or f"YouTube video {video_id}")[:200]
    except (httpx.HTTPError, ValueError):
        pass
    return f"YouTube video {video_id}"


async def _command(*args: str, timeout: float = 180) -> str:
    proc = await asyncio.create_subprocess_exec(*args, stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.STDOUT)
    try:
        output, _ = await asyncio.wait_for(proc.communicate(), timeout)
    except (asyncio.CancelledError, asyncio.TimeoutError):
        if proc.returncode is None:
            proc.kill()
        await proc.wait()
        raise
    if proc.returncode:
        message = output.decode(errors="replace")[-1200:]
        raise SourceError("YouTube could not provide this video's audio. It may be private, restricted, or temporarily blocked. " + message)
    return output.decode(errors="replace")


async def _transcribe_youtube(url: str, key: str) -> tuple[str, str, float]:
    """Download public video audio only and transcribe with the configured provider."""
    with tempfile.TemporaryDirectory(prefix="eduvid-youtube-", dir=SOURCES_DIR) as temporary:
        directory = Path(temporary)
        downloader = [sys.executable, "-m", "yt_dlp"]
        raw = await _command(*downloader, "--ignore-config", "--no-playlist", "--skip-download", "--dump-single-json",
                             "--no-warnings", "--socket-timeout", "20", "--retries", "1", url, timeout=90)
        info = json.loads(raw)
        duration = float(info.get("duration") or 0)
        if info.get("is_live") or not 0 < duration <= MAX_VIDEO_SECONDS:
            raise SourceError("Choose a recorded YouTube video of up to two hours.")
        await _command(*downloader, "--ignore-config", "--no-playlist", "--no-progress", "--no-warnings", "--no-cache-dir",
                       "--socket-timeout", "20", "--retries", "1", "--fragment-retries", "1", "--max-filesize", "80M",
                       "-f", "worstaudio", "-o", str(directory / "audio.%(ext)s"), url, timeout=300)
        candidates = [path for path in directory.glob("audio.*") if path.suffix not in (".part", ".ytdl")]
        if not candidates or candidates[0].stat().st_size > 80 * 1024 * 1024:
            raise SourceError("This video's audio is too large. Choose a shorter video.")
        audio = directory / "speech.mp3"
        await _command(shutil.which("ffmpeg") or "ffmpeg", "-y", "-v", "error", "-i", str(candidates[0]),
                       "-vn", "-ac", "1", "-ar", "16000", "-b:a", "32k", str(audio), timeout=120)
        async with httpx.AsyncClient(timeout=httpx.Timeout(300, connect=20)) as client:
            with audio.open("rb") as stream:
                response = await client.post("https://api.elevenlabs.io/v1/speech-to-text", headers={"xi-api-key": key},
                                             data={"model_id": "scribe_v2", "tag_audio_events": "false", "diarize": "false"},
                                             files={"file": ("speech.mp3", stream, "audio/mpeg")})
        if not response.is_success:
            from .elevenlabs import _explain
            raise SourceError("YouTube audio transcription failed. The ElevenLabs key needs Speech to Text access. " + _explain(response.status_code, response.text))
        text = clean_text(str(response.json().get("text") or ""))
        return text, str(info.get("title") or "YouTube video")[:200], duration


async def extract_youtube(url: str, elevenlabs_key: str | None = None) -> dict:
    video_id = youtube_video_id(url)
    canonical = f"https://www.youtube.com/watch?v={video_id}"
    warnings: list[str] = ["The explainer uses spoken content; details shown only in the video image are not read."]
    async with _YOUTUBE_LOCK:
        try:
            text, duration, generated, language = await asyncio.wait_for(asyncio.to_thread(_fetch_captions, video_id), timeout=65)
            title = await _youtube_title(canonical, video_id)
            method = "automatic captions" if generated else "captions"
            if generated:
                warnings.append("YouTube captions are automatic and may contain transcription errors.")
            if language != "English":
                warnings.append(f"Source captions are in {language}; narration uses your selected output language.")
        except asyncio.CancelledError:
            raise
        except Exception as e:
            if not elevenlabs_key or not elevenlabs_key.strip():
                raise SourceError("YouTube captions could not be retrieved. Add an ElevenLabs key with Speech to Text access for audio transcription, or upload a transcript.") from e
            try:
                text, title, duration = await _transcribe_youtube(canonical, elevenlabs_key.strip())
            except asyncio.TimeoutError as timeout_error:
                raise SourceError("YouTube audio retrieval timed out. Try a shorter video or upload its transcript.") from timeout_error
            except SourceError:
                raise
            except (httpx.HTTPError, ValueError, OSError) as error:
                raise SourceError("YouTube audio transcription could not be completed. Try again or upload a transcript.") from error
            method = "audio transcription"
            warnings.append("Captions could not be retrieved; audio was transcribed with ElevenLabs Scribe. Transcription uses provider credits.")
    if not isinstance(duration, (int, float)) or not math.isfinite(duration) or not 0 < duration <= MAX_VIDEO_SECONDS:
        raise SourceError("Choose a recorded YouTube video of up to two hours, or upload a selected transcript section.")
    if len(text) < 80:
        raise SourceError("This video has too little readable speech to make an explainer.")
    if len(text) > MAX_SOURCE_CHARS:
        raise SourceError("The transcript is too long. Upload a selected transcript section (up to 180,000 characters).")
    return {"type": "youtube", "title": title, "text": text, "url": canonical, "duration": round(duration, 2),
            "transcript_method": method, "warnings": warnings}


async def _public_article_url(url: str) -> str:
    """Reject local-network targets, including redirects, before sending a request."""
    url = url.strip()
    try:
        parsed = urlparse(url)
        if (parsed.scheme not in ("http", "https") or not parsed.hostname or parsed.username
                or parsed.password or parsed.port not in (None, 80, 443) or len(url) > 4000):
            raise ValueError("invalid public URL")
        addresses = await asyncio.get_running_loop().getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM)
        if not addresses or any(not ipaddress.ip_address(record[4][0]).is_global for record in addresses):
            raise ValueError("private address")
    except (ValueError, OSError) as e:
        raise SourceError("Use a public http or https article or PDF link.") from e
    return parsed._replace(fragment="").geturl()


async def extract_article_url(url: str) -> dict:
    original = url.strip()
    target = await _public_article_url(original)
    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(30, connect=15), follow_redirects=False,
                                     headers={"User-Agent": "EduVid/1.0 (article reader)", "Accept": "text/html,application/pdf,text/plain"}) as client:
            for _ in range(6):
                async with client.stream("GET", target) as response:
                    if response.is_redirect:
                        location = response.headers.get("location")
                        if not location:
                            raise SourceError("This article link returned an invalid redirect.")
                        target = await _public_article_url(urljoin(target, location))
                        continue
                    if response.status_code in (401, 403, 429):
                        raise SourceError("This website blocks direct article access. Download the article or copy its readable text into a TXT file, then upload it here.")
                    if not response.is_success:
                        raise SourceError(f"This article could not be retrieved (HTTP {response.status_code}). Upload a copy of the document instead.")
                    data = bytearray()
                    async for chunk in response.aiter_bytes():
                        data.extend(chunk)
                        if len(data) > MAX_UPLOAD_BYTES:
                            raise SourceError("This article is larger than 20 MB. Upload a smaller section.", 413)
                    content_type = response.headers.get("content-type", "").lower()
                    if bytes(data).lstrip().startswith(b"%PDF-") or "application/pdf" in content_type:
                        name = Path(urlparse(target).path).name or "article.pdf"
                        name = name if name.lower().endswith(".pdf") else name + ".pdf"
                    elif "text/html" in content_type or bytes(data).lstrip().lower().startswith((b"<!doctype html", b"<html")):
                        name = "article.html"
                        if any(marker in bytes(data).lower() for marker in (b"cf-chl-", b"challenge-platform")):
                            raise SourceError("This website requires browser verification. Download the article or copy its readable text into a TXT file, then upload it here.")
                    elif "text/plain" in content_type:
                        name = "article.txt"
                    else:
                        raise SourceError("This link is not a readable article or PDF. Upload a PDF, Word document, or text copy.")
                    source = await asyncio.to_thread(extract_article, bytes(data), name)
                    source.update(url=original, filename=None)
                    source["warnings"].append("Only article text is read; embedded images, diagrams and linked pages are not included.")
                    return source
            raise SourceError("This article redirects too many times. Use its final public link or upload a copy.")
    except httpx.TimeoutException as e:
        raise SourceError("Article retrieval timed out. Try again or upload a copy.") from e
    except httpx.HTTPError as e:
        raise SourceError("This article could not be reached. Check the link or upload a copy.") from e


def save_source(source: dict) -> dict:
    source = dict(source)
    source.update(id=uuid.uuid4().hex, created_at=time.time(), text_chars=len(source["text"]),
                  word_count=len(source["text"].split()), excerpt=source["text"][:1600])
    path = SOURCES_DIR / f"{source['id']}.json"
    path.write_text(json.dumps(source, ensure_ascii=False))
    return source


def get_source(source_id: str) -> dict:
    if not re.fullmatch(r"[a-f0-9]{32}", source_id or ""):
        raise SourceError("Invalid source reference. Extract your source again.")
    path = SOURCES_DIR / f"{source_id}.json"
    if not path.is_file():
        raise SourceError("Source not found. Upload the article or extract the video again.", 404)
    try:
        source = json.loads(path.read_text())
        if not isinstance(source, dict) or not isinstance(source.get("text"), str) or len(source["text"]) < 80:
            raise ValueError("invalid source text")
        return source
    except (ValueError, OSError) as e:
        raise SourceError("This source could not be read. Extract it again.") from e


def public_source(source: dict) -> dict:
    return {key: value for key, value in source.items() if key != "text"}


@router.post("/api/sources/article")
async def article_source(file: UploadFile = File(...)):
    try:
        data = await file.read(MAX_UPLOAD_BYTES + 1)
        if not data:
            raise SourceError("Choose a document containing readable text.")
        if len(data) > MAX_UPLOAD_BYTES:
            raise SourceError("Upload a document smaller than 20 MB.", 413)
        source = await asyncio.to_thread(extract_article, data, file.filename or "article.txt")
        return public_source(save_source(source))
    except SourceError as e:
        raise HTTPException(e.status_code, str(e)) from e
    finally:
        await file.close()


@router.post("/api/sources/youtube")
async def youtube_source(body: YouTubeSource):
    try:
        return public_source(save_source(await extract_youtube(body.url, body.elevenlabs_key)))
    except SourceError as e:
        raise HTTPException(e.status_code, str(e)) from e


@router.post("/api/sources/article-url")
async def article_url_source(body: ArticleURLSource):
    try:
        return public_source(save_source(await extract_article_url(body.url)))
    except SourceError as e:
        raise HTTPException(e.status_code, str(e)) from e


@router.get("/api/sources/{source_id}")
async def source_info(source_id: str):
    try:
        return public_source(get_source(source_id))
    except SourceError as e:
        raise HTTPException(e.status_code, str(e)) from e

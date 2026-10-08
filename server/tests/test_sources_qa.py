"""Offline QA for document ingestion, source APIs and YouTube caption fallback."""
from __future__ import annotations

import asyncio
import io
import json
import socket
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch

import httpx
from fastapi import FastAPI

from app import sources

# Scratch space for per-test temp dirs, inside the repo and git-ignored.
WORK = Path(__file__).resolve().parent / ".work"
WORK.mkdir(exist_ok=True)
VIDEO_ID = "dQw4w9WgXcQ"
CANONICAL = f"https://www.youtube.com/watch?v={VIDEO_ID}"
ARTICLE_TEXT = (
    "Safety brief: Inspect brakes, steering and tires before operating the vehicle. "
    "Keep personnel outside the work area. Wear the required protective equipment and follow the operator manual."
)


def pdf_document(text: str | None = ARTICLE_TEXT, *, encrypted: bool = False) -> bytes:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if text is not None:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        content.set_data(f"BT /F1 12 Tf 50 740 Td ({escaped}) Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)
    writer.add_metadata({"/Title": "Vehicle safety paper"})
    if encrypted:
        writer.encrypt("test password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


class SourceParsingQA(unittest.TestCase):
    def test_text_and_markdown_are_utf8_and_filename_is_sanitized(self):
        for filename in ("../folder/paper.txt", "C:\\folder\\paper.md"):
            with self.subTest(filename=filename):
                source = sources.extract_article(("\ufeff" + ARTICLE_TEXT).encode(), filename)
                self.assertEqual(source["text"], ARTICLE_TEXT)
                self.assertEqual(source["filename"], Path(filename.replace("\\", "/")).name)
                self.assertEqual(source["title"], "paper")

    def test_html_reads_article_text_without_executing_script_or_style(self):
        html = f"<html><head><title>Safety &amp; inspection</title><style>HIDDEN_CSS</style></head>" \
               f"<body><script>HIDDEN_SCRIPT</script><noscript>HIDDEN_FALLBACK</noscript><h1>Safety</h1><p>{ARTICLE_TEXT}</p></body></html>"
        source = sources.extract_article(html.encode(), "brief.html")
        self.assertEqual(source["title"], "Safety & inspection")
        self.assertIn("Inspect brakes", source["text"])
        self.assertNotIn("HIDDEN", source["text"])

    def test_selectable_pdf_extracts_text_and_warns_about_unread_diagrams(self):
        source = sources.extract_article(pdf_document(), "safety.pdf")
        self.assertIn("Inspect brakes", source["text"])
        self.assertEqual(source["title"], "Vehicle safety paper")
        self.assertTrue(any("diagrams and images" in warning for warning in source["warnings"]))

    def test_encrypted_pdf_has_clear_unlock_instruction(self):
        with self.assertRaisesRegex(sources.SourceError, "password protected"):
            sources.extract_article(pdf_document(encrypted=True), "protected.pdf")

    def test_scanned_or_empty_pdf_has_clear_ocr_instruction(self):
        with self.assertRaisesRegex(sources.SourceError, "Scanned PDFs need OCR"):
            sources.extract_article(pdf_document(None), "scanned.pdf")

    def test_docx_extracts_paragraphs_and_table_values(self):
        from docx import Document
        document = Document()
        document.add_paragraph(ARTICLE_TEXT)
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Maximum incline"
        table.cell(0, 1).text = "Consult the operator manual"
        output = io.BytesIO()
        document.save(output)
        source = sources.extract_article(output.getvalue(), "safety.docx")
        self.assertIn("Maximum incline | Consult the operator manual", source["text"])
        self.assertTrue(source["warnings"])

    def test_bad_files_fail_with_readable_errors(self):
        cases = (
            (b"not a PDF", "paper.pdf", "not a valid PDF"),
            (b"not a zip", "paper.docx", "could not be read"),
            (b"\xff" * 100, "paper.txt", "UTF-8"),
            (b"tiny", "paper.md", "too little readable text"),
            (ARTICLE_TEXT.encode(), "paper.exe", "Upload a PDF"),
        )
        for data, filename, error in cases:
            with self.subTest(filename=filename), self.assertRaisesRegex(sources.SourceError, error):
                sources.extract_article(data, filename)

    def test_docx_zip_without_word_document_is_rejected(self):
        output = io.BytesIO()
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("unrelated.txt", ARTICLE_TEXT)
        with self.assertRaisesRegex(sources.SourceError, "not a valid Word document"):
            sources.extract_article(output.getvalue(), "fake.docx")

    def test_source_text_length_limit_is_enforced(self):
        with patch.object(sources, "MAX_SOURCE_CHARS", 100):
            with self.assertRaisesRegex(sources.SourceError, "too long"):
                sources.extract_article(ARTICLE_TEXT.encode(), "long.txt")

    def test_clean_text_preserves_unicode_and_removes_control_characters(self):
        cleaned = sources.clean_text(" Café &amp; attention\x00\x01\r\n\r\n\r\n  安全 \t review ")
        self.assertEqual(cleaned, "Café & attention \n\n 安全 review")

    def test_youtube_link_variants_resolve_to_one_video_id(self):
        for url in (f"https://www.youtube.com/watch?v={VIDEO_ID}&t=20", f"youtu.be/{VIDEO_ID}",
                    f"https://m.youtube.com/shorts/{VIDEO_ID}", f"https://youtube.com/embed/{VIDEO_ID}",
                    f"https://youtube.com/live/{VIDEO_ID}"):
            with self.subTest(url=url):
                self.assertEqual(sources.youtube_video_id(url), VIDEO_ID)

    def test_non_video_and_untrusted_youtube_links_are_rejected(self):
        for url in ("", "https://example.com/watch?v=" + VIDEO_ID,
                    "https://www.youtube.com.evil.example/watch?v=" + VIDEO_ID,
                    "https://user:pass@www.youtube.com/watch?v=" + VIDEO_ID,
                    "https://www.youtube.com:444/watch?v=" + VIDEO_ID,
                    "https://www.youtube.com/playlist?list=abc", "https://youtu.be/short"):
            with self.subTest(url=url), self.assertRaises(sources.SourceError):
                sources.youtube_video_id(url)

    def test_malformed_youtube_urls_have_readable_errors(self):
        for url in ("https://www.youtube.com:bad/watch?v=" + VIDEO_ID,
                    "https://www.youtube.com:99999/watch?v=" + VIDEO_ID,
                    "https://[www.youtube.com/watch?v=" + VIDEO_ID):
            with self.subTest(url=url), self.assertRaises(sources.SourceError):
                sources.youtube_video_id(url)

    def test_caption_parser_prefers_manual_english_and_preserves_original_duration(self):
        class Fetched(list):
            language = "English"
            is_generated = False
        fetched = Fetched([SimpleNamespace(text="Inspect &amp; verify brakes.", start=2, duration=3),
                           SimpleNamespace(text=ARTICLE_TEXT, start=17, duration=4)])
        manual_en = Mock(language_code="en", is_generated=False)
        manual_en.fetch.return_value = fetched
        auto_en = Mock(language_code="en-US", is_generated=True)
        manual_fr = Mock(language_code="fr", is_generated=False)
        api = Mock()
        api.list.return_value = [manual_fr, auto_en, manual_en]
        with patch("youtube_transcript_api.YouTubeTranscriptApi", return_value=api):
            text, duration, generated, language = sources._fetch_captions(VIDEO_ID)
        self.assertIn("Inspect & verify brakes.", text)
        self.assertEqual(duration, 21)
        self.assertFalse(generated)
        self.assertEqual(language, "English")
        manual_en.fetch.assert_called_once()
        auto_en.fetch.assert_not_called()
        manual_fr.fetch.assert_not_called()

    def test_caption_parser_empty_list_has_readable_error(self):
        api = Mock()
        api.list.return_value = []
        with patch("youtube_transcript_api.YouTubeTranscriptApi", return_value=api):
            with self.assertRaisesRegex(sources.SourceError, "No captions"):
                sources._fetch_captions(VIDEO_ID)


class SourceAPIQA(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        WORK.mkdir(parents=True, exist_ok=True)
        self.tmp = tempfile.TemporaryDirectory(dir=WORK)
        self.directory = Path(self.tmp.name)
        self.directory_patch = patch.object(sources, "SOURCES_DIR", self.directory)
        self.directory_patch.start()
        exists = patch.object(sources, "_youtube_video_missing", AsyncMock(return_value=False))
        exists.start()
        self.addCleanup(exists.stop)
        self.app = FastAPI()
        self.app.include_router(sources.router)
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=self.app, raise_app_exceptions=False),
                                       base_url="http://test")

    async def asyncTearDown(self):
        await self.client.aclose()
        self.directory_patch.stop()
        self.tmp.cleanup()

    async def test_article_upload_returns_review_metadata_and_persists_full_text_locally(self):
        response = await self.client.post("/api/sources/article", files={"file": ("brief.txt", ARTICLE_TEXT.encode(), "text/plain")})
        self.assertEqual(response.status_code, 200)
        metadata = response.json()
        self.assertNotIn("text", metadata)
        self.assertEqual(metadata["text_chars"], len(ARTICLE_TEXT))
        self.assertEqual(metadata["word_count"], len(ARTICLE_TEXT.split()))
        self.assertRegex(metadata["id"], r"^[a-f0-9]{32}$")
        self.assertEqual(sources.get_source(metadata["id"])["text"], ARTICLE_TEXT)
        reviewed = await self.client.get("/api/sources/" + metadata["id"])
        self.assertEqual(reviewed.status_code, 200)
        self.assertNotIn("text", reviewed.json())

    async def test_upload_empty_unsupported_and_oversize_errors_are_json_and_create_no_source(self):
        for content, filename, expected in ((b"", "empty.txt", 400), (ARTICLE_TEXT.encode(), "bad.exe", 400)):
            with self.subTest(filename=filename):
                response = await self.client.post("/api/sources/article", files={"file": (filename, content)})
                self.assertEqual(response.status_code, expected)
                self.assertIsInstance(response.json()["detail"], str)
        with patch.object(sources, "MAX_UPLOAD_BYTES", 10):
            response = await self.client.post("/api/sources/article", files={"file": ("large.txt", ARTICLE_TEXT.encode())})
            self.assertEqual(response.status_code, 413)
            self.assertIn("20 MB", response.json()["detail"])
        self.assertEqual(list(self.directory.iterdir()), [])

    async def test_source_reference_errors_are_readable(self):
        response = await self.client.get("/api/sources/" + "a" * 32)
        self.assertEqual(response.status_code, 404)
        self.assertIn("Source not found", response.json()["detail"])
        response = await self.client.get("/api/sources/bad-reference")
        self.assertEqual(response.status_code, 400)

    async def test_corrupt_stored_source_json_is_not_a_server_error(self):
        source_id = "b" * 32
        for payload in ("{broken", "[]", "null", '{"text": "too short"}'):
            with self.subTest(payload=payload):
                (self.directory / f"{source_id}.json").write_text(payload)
                response = await self.client.get("/api/sources/" + source_id)
                self.assertEqual(response.status_code, 400)
                self.assertIn("could not be read", response.json()["detail"])

    async def test_invalid_youtube_api_input_does_not_return_500(self):
        for url in ("https://www.youtube.com:bad/watch?v=" + VIDEO_ID, "https://[www.youtube.com/watch?v=" + VIDEO_ID):
            with self.subTest(url=url):
                response = await self.client.post("/api/sources/youtube", json={"url": url})
                self.assertEqual(response.status_code, 400)
                self.assertIsInstance(response.json()["detail"], str)

    async def test_youtube_api_saves_caption_source_without_persisting_provider_key(self):
        with patch.object(sources, "_fetch_captions", return_value=(ARTICLE_TEXT, 42.0, False, "English")), \
             patch.object(sources, "_youtube_title", AsyncMock(return_value="Inspection training")), \
             patch.object(sources, "_transcribe_youtube", AsyncMock()) as transcription:
            response = await self.client.post("/api/sources/youtube", json={"url": "https://youtu.be/" + VIDEO_ID,
                                                                        "elevenlabs_key": "private key marker"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["transcript_method"], "captions")
        self.assertEqual(response.json()["url"], CANONICAL)
        self.assertNotIn("text", response.json())
        transcription.assert_not_awaited()
        persisted = (self.directory / (response.json()["id"] + ".json")).read_text()
        self.assertNotIn("private key marker", persisted)

    async def test_caption_fallback_timeout_returns_readable_json_error(self):
        with patch.object(sources, "_fetch_captions", side_effect=sources.SourceError("No captions")), \
             patch.object(sources, "_transcribe_youtube", AsyncMock(side_effect=asyncio.TimeoutError)):
            response = await self.client.post("/api/sources/youtube", json={"url": CANONICAL, "elevenlabs_key": "unused"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("timed out", response.json()["detail"])

    async def test_caption_fallback_network_error_returns_readable_json_error(self):
        with patch.object(sources, "_fetch_captions", side_effect=sources.SourceError("No captions")), \
             patch.object(sources, "_transcribe_youtube", AsyncMock(side_effect=httpx.ConnectError("offline"))):
            response = await self.client.post("/api/sources/youtube", json={"url": CANONICAL, "elevenlabs_key": "unused"})
        self.assertEqual(response.status_code, 400)
        self.assertIn("transcription", response.json()["detail"].lower())

    def public_dns(self):
        return patch.object(asyncio.get_running_loop(), "getaddrinfo", AsyncMock(return_value=[
            (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443)),
        ]))

    def remote_http(self, handler):
        original = httpx.AsyncClient
        return patch.object(httpx, "AsyncClient", side_effect=lambda **kwargs: original(
            **kwargs, transport=httpx.MockTransport(handler)))

    async def test_article_url_extracts_main_text_and_preserves_original_reference(self):
        url = "https://manualmachine.com/cat/ad60/9684589-brochure/#10"
        html = f"<html><head><title>AD60 safety brochure</title></head><body><nav>NAVIGATION_NOISE</nav>" \
               f"<main><h1>Vehicle safety</h1><p>{ARTICLE_TEXT}</p></main><footer>FOOTER_NOISE</footer></body></html>"
        requested = []
        def respond(request):
            requested.append(str(request.url))
            return httpx.Response(200, text=html, headers={"content-type": "text/html"})
        with self.public_dns(), self.remote_http(respond):
            response = await self.client.post("/api/sources/article-url", json={"url": url})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["title"], "AD60 safety brochure")
        self.assertEqual(response.json()["url"], url)
        self.assertNotIn("text", response.json())
        source = sources.get_source(response.json()["id"])
        self.assertIn("Inspect brakes", source["text"])
        self.assertNotIn("NOISE", source["text"])
        self.assertEqual(requested, [url.split("#")[0]])
        self.assertTrue(any("embedded images" in warning for warning in source["warnings"]))

    async def test_article_url_can_extract_pdf_without_pdf_filename_extension(self):
        with self.public_dns(), self.remote_http(lambda request: httpx.Response(
                200, content=pdf_document(), headers={"content-type": "application/pdf"})):
            source = await sources.extract_article_url("https://example.com/download?id=123")
        self.assertEqual(source["title"], "Vehicle safety paper")
        self.assertIn("Inspect brakes", source["text"])
        self.assertIsNone(source["filename"])

    async def test_article_url_revalidates_redirect_target_before_fetching_private_address(self):
        requested = []
        async def resolve(host, port, **kwargs):
            address = "127.0.0.1" if host == "127.0.0.1" else "93.184.216.34"
            return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (address, port))]
        def respond(request):
            requested.append(str(request.url))
            return httpx.Response(302, headers={"location": "http://127.0.0.1/private.txt"})
        with patch.object(asyncio.get_running_loop(), "getaddrinfo", resolve), self.remote_http(respond):
            with self.assertRaisesRegex(sources.SourceError, "public http or https"):
                await sources.extract_article_url("https://example.com/redirect")
        self.assertEqual(requested, ["https://example.com/redirect"])

    async def test_article_url_rejects_credentials_schemes_ports_and_private_dns(self):
        for url in ("file:///etc/passwd", "ftp://example.com/paper", "https://user:pass@example.com/paper",
                    "https://example.com:bad/paper", "https://example.com:444/paper"):
            with self.subTest(url=url), self.assertRaises(sources.SourceError):
                await sources._public_article_url(url)
        with patch.object(asyncio.get_running_loop(), "getaddrinfo", AsyncMock(return_value=[
            (socket.AF_INET, socket.SOCK_STREAM, 6, "", ("10.0.0.4", 443)),
        ])):
            with self.assertRaises(sources.SourceError):
                await sources._public_article_url("https://internal.example/paper")

    async def test_article_url_public_redirect_chain_preserves_original_reference(self):
        original_url = "https://example.com/redirect#section"
        requested = []
        def respond(request):
            requested.append(str(request.url))
            if request.url.path == "/redirect":
                return httpx.Response(302, headers={"location": "/paper.txt"})
            return httpx.Response(200, text=ARTICLE_TEXT, headers={"content-type": "text/plain"})
        with self.public_dns() as dns, self.remote_http(respond):
            source = await sources.extract_article_url(original_url)
        self.assertEqual(source["url"], original_url)
        self.assertEqual(requested, ["https://example.com/redirect", "https://example.com/paper.txt"])
        self.assertEqual(dns.await_count, 2)

    async def test_blocked_or_challenge_article_url_explains_upload_alternative(self):
        for status, html in ((403, "Forbidden"), (200, "<html><p>challenge-platform: browser verification</p></html>")):
            with self.subTest(status=status), self.public_dns(), self.remote_http(lambda request: httpx.Response(
                    status, text=html, headers={"content-type": "text/html"})):
                response = await self.client.post("/api/sources/article-url", json={"url": "https://example.com/article"})
                self.assertEqual(response.status_code, 400)
                self.assertIn("TXT file", response.json()["detail"])
        self.assertEqual(list(self.directory.iterdir()), [])

    async def test_article_url_download_size_is_limited_before_parsing(self):
        with patch.object(sources, "MAX_UPLOAD_BYTES", 100), self.public_dns(), self.remote_http(
                lambda request: httpx.Response(200, text=ARTICLE_TEXT, headers={"content-type": "text/plain"})):
            response = await self.client.post("/api/sources/article-url", json={"url": "https://example.com/large.txt"})
        self.assertEqual(response.status_code, 413)
        self.assertIn("larger than", response.json()["detail"])
        self.assertEqual(list(self.directory.iterdir()), [])

    async def test_article_url_redirects_are_bounded(self):
        requested = []
        def respond(request):
            requested.append(str(request.url))
            return httpx.Response(302, headers={"location": "/loop"})
        with self.public_dns(), self.remote_http(respond):
            with self.assertRaisesRegex(sources.SourceError, "too many times"):
                await sources.extract_article_url("https://example.com/loop")
        self.assertLessEqual(len(requested), 6)

    async def test_article_url_timeout_and_unreachable_errors_have_readable_json(self):
        for exception, wording in ((httpx.ReadTimeout("timeout"), "timed out"), (httpx.ConnectError("offline"), "could not be reached")):
            def fail(request):
                raise exception
            with self.subTest(wording=wording), self.public_dns(), self.remote_http(fail):
                response = await self.client.post("/api/sources/article-url", json={"url": "https://example.com/article"})
                self.assertEqual(response.status_code, 400)
                self.assertIn(wording, response.json()["detail"])

    async def test_article_url_rejects_non_document_content(self):
        with self.public_dns(), self.remote_http(lambda request: httpx.Response(
                200, content=b"binary archive", headers={"content-type": "application/zip"})):
            with self.assertRaisesRegex(sources.SourceError, "not a readable article or PDF"):
                await sources.extract_article_url("https://example.com/archive")

    async def fake_audio_command(self, *args, timeout):
        if "--dump-single-json" in args:
            return json.dumps({"duration": 90, "title": "Safety training"})
        if "-o" in args:
            Path(args[args.index("-o") + 1].replace("%(ext)s", "m4a")).write_bytes(b"fixture audio")
        else:
            Path(args[-1]).write_bytes(b"fixture speech mp3")
        return ""

    async def test_full_transcription_fallback_uploads_audio_and_cleans_up_temporary_files(self):
        requests = []
        def respond(request):
            requests.append(request)
            return httpx.Response(200, json={"text": ARTICLE_TEXT})
        with patch.object(sources, "_command", self.fake_audio_command), self.remote_http(respond):
            text, title, duration = await sources._transcribe_youtube(CANONICAL, "unused private key")
        self.assertEqual(text, ARTICLE_TEXT)
        self.assertEqual(title, "Safety training")
        self.assertEqual(duration, 90)
        self.assertEqual(len(requests), 1)
        self.assertEqual(str(requests[0].url), "https://api.elevenlabs.io/v1/speech-to-text")
        self.assertEqual(requests[0].headers["xi-api-key"], "unused private key")
        self.assertIn(b"scribe_v2", requests[0].content)
        self.assertIn(b"fixture speech mp3", requests[0].content)
        self.assertEqual(list(self.directory.iterdir()), [])

    async def test_transcription_permission_error_is_readable_and_temp_files_are_removed(self):
        with patch.object(sources, "_command", self.fake_audio_command), self.remote_http(lambda request: httpx.Response(
                401, json={"detail": {"status": "missing_permissions", "message": "Missing Speech to Text permission"}})):
            with self.assertRaisesRegex(sources.SourceError, "Speech to Text access"):
                await sources._transcribe_youtube(CANONICAL, "unused")
        self.assertEqual(list(self.directory.iterdir()), [])

    async def test_transcription_rejects_live_or_overlong_videos_before_downloading_audio(self):
        for info in ({"duration": 90, "is_live": True}, {"duration": 7201}, {"duration": 0}):
            command = AsyncMock(return_value=json.dumps(info))
            with self.subTest(info=info), patch.object(sources, "_command", command):
                with self.assertRaisesRegex(sources.SourceError, "recorded YouTube video"):
                    await sources._transcribe_youtube(CANONICAL, "unused")
            self.assertEqual(command.await_count, 1)
        self.assertEqual(list(self.directory.iterdir()), [])


class YouTubeExtractionQA(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        exists = patch.object(sources, "_youtube_video_missing", AsyncMock(return_value=False))
        exists.start()
        self.addCleanup(exists.stop)

    async def test_missing_video_is_reported_without_caption_or_paid_fallback_advice(self):
        for key in (None, "unused"):
            with self.subTest(key=key), \
                 patch.object(sources, "_youtube_video_missing", AsyncMock(return_value=True)), \
                 patch.object(sources, "_fetch_captions", side_effect=sources.SourceError("No captions")), \
                 patch.object(sources, "_transcribe_youtube", AsyncMock()) as transcription:
                with self.assertRaisesRegex(sources.SourceError, "doesn't exist or was removed"):
                    await sources.extract_youtube(CANONICAL, key)
                transcription.assert_not_awaited()

    async def test_automatic_and_foreign_caption_warnings_are_preserved(self):
        with patch.object(sources, "_fetch_captions", return_value=(ARTICLE_TEXT, 75.0, True, "French")), \
             patch.object(sources, "_youtube_title", AsyncMock(return_value="Safety video")), \
             patch.object(sources, "_transcribe_youtube", AsyncMock()) as transcription:
            source = await sources.extract_youtube(CANONICAL, "unused")
        self.assertEqual(source["transcript_method"], "automatic captions")
        self.assertTrue(any("automatic" in warning for warning in source["warnings"]))
        self.assertTrue(any("French" in warning for warning in source["warnings"]))
        transcription.assert_not_awaited()

    async def test_unavailable_captions_without_key_guide_user_to_transcript_or_permission(self):
        with patch.object(sources, "_fetch_captions", side_effect=sources.SourceError("No captions")), \
             patch.object(sources, "_transcribe_youtube", AsyncMock()) as transcription:
            with self.assertRaisesRegex(sources.SourceError, "Speech to Text"):
                await sources.extract_youtube(CANONICAL)
        transcription.assert_not_awaited()

    async def test_unavailable_captions_use_one_transcription_fallback_and_trim_key(self):
        transcription = AsyncMock(return_value=(ARTICLE_TEXT, "Transcribed safety video", 80.0))
        with patch.object(sources, "_fetch_captions", side_effect=sources.SourceError("No captions")), \
             patch.object(sources, "_transcribe_youtube", transcription):
            source = await sources.extract_youtube("https://youtu.be/" + VIDEO_ID, "  unused  ")
        transcription.assert_awaited_once_with(CANONICAL, "unused")
        self.assertEqual(source["transcript_method"], "audio transcription")
        self.assertTrue(any("provider credits" in warning for warning in source["warnings"]))

    async def test_generic_caption_failure_also_uses_one_fallback(self):
        transcription = AsyncMock(return_value=(ARTICLE_TEXT, "Safety video", 80.0))
        with patch.object(sources, "_fetch_captions", side_effect=RuntimeError("Upstream blocked")), \
             patch.object(sources, "_transcribe_youtube", transcription):
            source = await sources.extract_youtube(CANONICAL, "unused")
        self.assertEqual(transcription.await_count, 1)
        self.assertEqual(source["transcript_method"], "audio transcription")

    async def test_overlong_and_empty_caption_sources_are_rejected_without_paid_fallback(self):
        for text, duration, error in ((ARTICLE_TEXT, 7201, "up to two hours"), ("tiny", 30, "too little readable speech")):
            with self.subTest(error=error), \
                 patch.object(sources, "_fetch_captions", return_value=(text, duration, False, "English")), \
                 patch.object(sources, "_youtube_title", AsyncMock(return_value="Safety video")), \
                 patch.object(sources, "_transcribe_youtube", AsyncMock()) as transcription:
                with self.assertRaisesRegex(sources.SourceError, error):
                    await sources.extract_youtube(CANONICAL, "unused")
                transcription.assert_not_awaited()

    async def test_caption_retrieval_cancellation_never_starts_paid_fallback(self):
        with patch.object(sources.asyncio, "to_thread", AsyncMock(side_effect=asyncio.CancelledError)), \
             patch.object(sources, "_transcribe_youtube", AsyncMock()) as transcription:
            with self.assertRaises(asyncio.CancelledError):
                await sources.extract_youtube(CANONICAL, "unused")
        transcription.assert_not_awaited()

    async def test_title_network_failure_keeps_valid_captions(self):
        original = httpx.AsyncClient
        def offline(request):
            raise httpx.ConnectError("offline", request=request)
        with patch.object(httpx, "AsyncClient", side_effect=lambda **kwargs: original(
                **kwargs, transport=httpx.MockTransport(offline))):
            self.assertEqual(await sources._youtube_title(CANONICAL, VIDEO_ID), f"YouTube video {VIDEO_ID}")

    async def test_malformed_title_metadata_does_not_discard_valid_captions(self):
        original = httpx.AsyncClient
        with patch.object(httpx, "AsyncClient", side_effect=lambda **kwargs: original(
                **kwargs, transport=httpx.MockTransport(lambda request: httpx.Response(200, json=[])))):
            self.assertEqual(await sources._youtube_title(CANONICAL, VIDEO_ID), f"YouTube video {VIDEO_ID}")

    async def test_invalid_caption_duration_has_readable_error(self):
        for duration in (float("nan"), -1, 0):
            with self.subTest(duration=duration), \
                 patch.object(sources, "_fetch_captions", return_value=(ARTICLE_TEXT, duration, False, "English")), \
                 patch.object(sources, "_youtube_title", AsyncMock(return_value="Safety video")):
                with self.assertRaises(sources.SourceError):
                    await sources.extract_youtube(CANONICAL)


if __name__ == "__main__":
    unittest.main()


class YouTubeExistenceQA(unittest.IsolatedAsyncioTestCase):
    async def test_only_an_oembed_404_counts_as_missing(self):
        original = httpx.AsyncClient
        def client_for(handler):
            return lambda **kwargs: original(**kwargs, transport=httpx.MockTransport(handler))
        def offline(request):
            raise httpx.ConnectError("offline", request=request)
        cases = {
            "not found": (lambda request: httpx.Response(404), True),
            "public": (lambda request: httpx.Response(200, json={"title": "x"}), False),
            "private or embed-disabled": (lambda request: httpx.Response(401), False),
            "offline": (offline, False),
        }
        for name, (handler, missing) in cases.items():
            with self.subTest(name), patch.object(httpx, "AsyncClient", side_effect=client_for(handler)):
                self.assertIs(await sources._youtube_video_missing(CANONICAL), missing)
